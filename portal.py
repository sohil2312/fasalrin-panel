#!/usr/bin/env python3
"""
Fasalrin control panel — local web page to run the scripts and watch them live.

    .venv\\Scripts\\python portal.py        -> opens http://localhost:8765

Runs on the branch PC only (bound to 127.0.0.1): CSV, Aadhaar and account numbers never leave it.
Starts fasalrin_regular.py (entry, Branch User login) or fasalrin_verify.py (approval, Branch Head
login) as a child process, answers their questions from the page, and parses their log for the
live view. "Stop after current row" creates stop.flag, which both scripts check between rows.
"""

from __future__ import annotations

import collections
import csv
import json
import os
import re
import subprocess
import sys
import threading
import time
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import branches

ROOT = Path(__file__).resolve().parent
HOST = "127.0.0.1"
PORT = int(os.environ.get("FASALRIN_PORT") or 8765)   # the portable zip uses its own port (Start Fasalrin.bat)
# extra dashboards (another port: FASALRIN_PORT=8767 ...) keep their own stop flag, job logs and browser
# profiles, so different branches run side by side; one branch never runs on two dashboards at once
INSTANCE = "" if PORT in (8765, 8766) else str(PORT)     # 8766 = the portable zip's main dashboard
STOP_FLAG = ROOT / (f"stop_{INSTANCE}.flag" if INSTANCE else "stop.flag")
RUNNING_DIR = ROOT / ".running"                     # <port>.txt = the work list a dashboard's job runs on


def log_name(name: str) -> str:
    return f"{Path(name).stem}_{INSTANCE}.log" if INSTANCE else name


def sol_dir_of(csv_name: str) -> str:
    """branches/<SOL> of a work list (any scheme), '' for loose CSVs."""
    parts = Path(csv_name).parts
    return "/".join(parts[:2]) if len(parts) > 2 and parts[0] == "branches" else ""


def _pid_alive(pid: int) -> bool:
    if os.name == "nt":                               # (os.kill(pid, 0) would END the process on Windows)
        r = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/NH"], capture_output=True, text=True)
        return str(pid) in r.stdout
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


# approval jobs that may run on SEVERAL dashboards for the same branch at once (each reserves a claim before
# opening it, so two dashboards never take the same one); every other job keeps a branch to itself
SHARED_JOBS = {"pri_approve", "prireg_approve"}


def busy_elsewhere(csv_name: str, kind: str = "") -> str:
    """The port of another dashboard whose running job is on the same branch, '' if none (or when both jobs are
    the same shared approval job)."""
    me, sd = str(PORT), sol_dir_of(csv_name)
    for p in RUNNING_DIR.glob("*.txt") if RUNNING_DIR.exists() else []:
        if p.stem == me:
            continue
        try:
            pid_s, other, *rest = p.read_text(encoding="utf-8").split("\n")
        except (OSError, ValueError):
            continue
        other_kind = rest[0].strip() if rest else ""
        if not _pid_alive(int(pid_s)):
            p.unlink(missing_ok=True)                  # that job has ended
            continue
        if other.strip() == csv_name or (sd and sol_dir_of(other.strip()) == sd):
            if kind in SHARED_JOBS and other_kind == kind and other.strip() == csv_name:
                continue                               # the same approval on the same list: shared, allowed
            return p.stem
    return ""
JOBS = {
    "entry":   {"script": "fasalrin_regular.py", "log": "run.log",    "login": "Branch User"},
    "entry_fresh": {"script": "fasalrin_fresh.py", "log": "fresh.log", "login": "Branch User"},
    "approve": {"script": "fasalrin_verify.py",  "log": "verify.log", "login": "Branch Head"},
    "claim":   {"script": "fasalrin_claim.py",   "log": "claim.log",  "login": "Branch User"},
    "claim_approve": {"script": "fasalrin_claim_verify.py", "log": "claim_verify.log", "login": "Branch Head"},
    # IS additional (rollover claims for 2024-25 loans)
    "additional": {"script": "fasalrin_additional.py", "log": "additional.log", "login": "Branch User"},
    "additional_approve": {"script": "fasalrin_additional_verify.py", "log": "additional_verify.log", "login": "Branch Head"},
    # PRI additional (3% PRI claims for 2024-25 loans)
    "pri": {"script": "fasalrin_pri.py", "log": "pri.log", "login": "Branch User"},
    "pri_approve": {"script": "fasalrin_pri_verify.py", "log": "pri_verify.log", "login": "Branch Head"},
    # PRI regular (3% PRI claims for 2025-26 loans; claim only, the loan is entered through IS regular)
    "prireg": {"script": "fasalrin_prireg.py", "log": "prireg.log", "login": "Branch User"},
    "prireg_approve": {"script": "fasalrin_prireg_verify.py", "log": "prireg_verify.log", "login": "Branch Head"},
    # IS fresh (farmers not in the portal yet: entered from scratch from the "not in system" Excel)
    "fresh": {"script": "fasalrin_fresh.py", "log": "fresh.log", "login": "Branch User"},
    # IS regular claim (1.5% IS on 2025-26 loans) from its own claim file: every account of the branch
    "isclaim": {"script": "fasalrin_claim.py", "log": "isclaim.log", "login": "Branch User"},
    "isclaim_approve": {"script": "fasalrin_claim_verify.py", "log": "isclaim_verify.log", "login": "Branch Head"},
}
ADDITIONAL_JOBS = {"additional", "additional_approve"}
JOB_SCHEME = {"additional": "additional", "additional_approve": "additional", "pri": "pri", "pri_approve": "pri",
              "prireg": "prireg", "prireg_approve": "prireg", "fresh": "fresh",
              "isclaim": "isclaim", "isclaim_approve": "isclaim"}
SCHEME_NAME = {"regular": "Loan Applications", "additional": "IS Additional Claims", "pri": "PRI additional", "prireg": "PRI regular",
               "fresh": "IS fresh", "isclaim": "IS Regular Claims"}
MASTER_KINDS = ("regular", "additional", "pri", "prireg", "fresh", "isclaim")
SIDE_FILES = ("_progress.csv", "_approvals.csv")

lock = threading.Lock()
job: dict = {}                       # the one job this panel runs at a time


# ----------------------------------------------------------------------------
# jobs
# ----------------------------------------------------------------------------
def input_csvs() -> list[str]:
    """Branch work lists (branches/<SOL>/<SOL>.csv) first, then loose CSVs in the folder."""
    work = sorted(p.relative_to(ROOT).as_posix() for p in (ROOT / "branches").glob("*/*.csv")
                  if p.stem == p.parent.name)
    add = sorted(p.relative_to(ROOT).as_posix() for p in (ROOT / "branches").glob("*/additional/*_additional.csv")
                 if p.stem == f"{p.parent.parent.name}_additional")
    pri = sorted(p.relative_to(ROOT).as_posix() for k in ("pri", "prireg")
                 for p in (ROOT / "branches").glob(f"*/{k}/*_{k}.csv") if p.stem == f"{p.parent.parent.name}_{k}")
    fresh = sorted(p.relative_to(ROOT).as_posix() for p in (ROOT / "branches").glob("*/fresh/*_fresh.csv")
                   if p.stem == f"{p.parent.parent.name}_fresh")
    isclaim = sorted(p.relative_to(ROOT).as_posix() for p in (ROOT / "branches").glob("*/isclaim/*_isclaim.csv")
                     if p.stem == f"{p.parent.parent.name}_isclaim")
    loose = sorted(p.name for p in ROOT.glob("*.csv") if not p.name.endswith(SIDE_FILES))
    return work + add + pri + fresh + isclaim + loose


MASTER_DIR = ROOT / "master"
MASTER_DIRS = {"regular": MASTER_DIR, "additional": MASTER_DIR / "additional", "pri": MASTER_DIR / "pri",
               "prireg": MASTER_DIR / "prireg", "fresh": MASTER_DIR / "fresh", "isclaim": MASTER_DIR / "isclaim"}
SOLS_OF = {"regular": branches.sols, "additional": branches.sols_additional, "pri": branches.sols_pri,
           "prireg": branches.sols_prireg, "fresh": branches.sols_fresh, "isclaim": branches.sols_isclaim}
BUILD_OF = {"regular": branches.build, "additional": branches.build_additional, "pri": branches.build_pri,
            "prireg": branches.build_prireg, "fresh": branches.build_fresh, "isclaim": branches.build_isclaim}


def masters(kind: str = "regular") -> list[str]:
    d = MASTER_DIRS.get(kind, MASTER_DIR)
    return sorted((p.name for p in d.glob("*.xlsx")), key=lambda n: (d / n).stat().st_mtime,
                  reverse=True) if d.exists() else []


def save_master(name: str, b64: str, kind: str = "regular") -> tuple[str | None, str | None]:
    import base64
    base = os.path.basename(name or "").strip()
    if not re.fullmatch(r"[\w .()%&+,-]{1,100}\.xlsx", base, re.I):
        return "master must be an .xlsx file", None
    try:
        data = base64.b64decode(b64, validate=True)
    except ValueError:
        return "bad file data", None
    if len(data) > MAX_UPLOAD or not data.startswith(b"PK"):
        return "not an Excel .xlsx file (or too large)", None
    d = MASTER_DIRS.get(kind, MASTER_DIR)
    d.mkdir(parents=True, exist_ok=True)
    dest = d / base
    if dest.exists():
        dest = d / f"{dest.stem}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    dest.write_bytes(data)
    try:                                          # a hand-work export is not a master
        import openpyxl
        wb = openpyxl.load_workbook(dest, read_only=True)
        head = [str(c or "").strip() for c in next(wb.worksheets[0].iter_rows(max_row=1, values_only=True), [])]
        wb.close()
        if "Fixed (Y)" in head or ("Reason" in head and "What to do" in head):
            dest.unlink(missing_ok=True)
            return ("this is a hand-work file (Reason / Fixed (Y)), not a master: upload it under "
                    "'Upload corrected CSVs' in CSV totals"), None
    except Exception:
        pass
    try:                                          # validates the columns for that master type
        SOLS_OF.get(kind, branches.sols)(dest)
    except Exception as e:
        dest.unlink(missing_ok=True)
        return str(e)[:200], None
    return None, dest.name


def xlsx_to_csv_text(b64: str) -> str:
    """Edited hand-work .xlsx -> CSV text (dates dd-mm-yyyy, whole numbers without .0 or E+)."""
    import base64
    import io
    from datetime import date as _d, datetime as _dt
    import openpyxl
    raw = base64.b64decode(b64, validate=True)
    if len(raw) > MAX_UPLOAD or not raw.startswith(b"PK"):
        raise ValueError("not an .xlsx file (or too large)")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = io.StringIO()
    w = csv.writer(out)
    for r in wb.worksheets[0].iter_rows(values_only=True):
        if not any(c not in (None, "") for c in r):
            continue
        w.writerow([v.strftime("%d-%m-%Y") if isinstance(v, (_dt, _d)) else
                    str(int(v)) if isinstance(v, float) and v.is_integer() else
                    "" if v is None else str(v).strip() for v in r])
    wb.close()
    return out.getvalue()


def csv_to_xlsx(name: str, b64: str) -> tuple[str, str]:
    """IS fresh upload as CSV -> (xlsx name, xlsx base64). Every cell stays text (long numbers intact)."""
    import base64
    import io
    import openpyxl
    raw = base64.b64decode(b64, validate=True)
    if len(raw) > MAX_UPLOAD:
        raise ValueError("file too large")
    text = raw.decode("utf-8-sig", errors="replace")
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise ValueError("the CSV has no data rows")
    head = [h.strip() for h in rows[0]]
    for col in ("Account No.", "Aadhaar No."):
        if col in head and any("E+" in (r[head.index(col)] if head.index(col) < len(r) else "").upper() for r in rows[1:]):
            raise ValueError(f"{col} looks like 9.99E+13 (Excel changed the numbers): format the column as Text and save again")
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append([c.strip() for c in r])
    buf = io.BytesIO()
    wb.save(buf)
    return os.path.splitext(os.path.basename(name))[0] + ".xlsx", base64.b64encode(buf.getvalue()).decode()


REQUIRED_COLS = ("Account No.", "Aadhaar No.", "Disb. Date", "DP")
MAX_UPLOAD = 20 * 1024 * 1024


def save_upload(name: str, text: str) -> tuple[str | None, str | None]:
    """Validate + save an uploaded CSV next to the scripts. Returns (error, saved name)."""
    base = os.path.basename(name or "").strip()
    if not re.fullmatch(r"[\w .()-]{1,80}\.csv", base, re.I) or base.endswith(SIDE_FILES):
        return "file name must be a plain .csv name", None
    if len(text.encode("utf-8")) > MAX_UPLOAD:
        return "file too large", None
    try:
        header = next(csv.reader([text.lstrip("﻿").splitlines()[0]]))
    except (StopIteration, IndexError, csv.Error):
        return "file is empty or not a CSV", None
    missing = [c for c in REQUIRED_COLS if c not in [h.strip() for h in header]]
    if missing:
        return "missing columns: " + ", ".join(missing), None
    if (ROOT / base).exists():                   # never overwrite: its Status column is the progress record
        return f"{base} already exists here — rename the file to upload it as a new list", None
    (ROOT / base).write_text(text, encoding="utf-8", newline="")
    return None, base


def running() -> bool:
    return bool(job) and job["proc"].poll() is None


def start_job(kind: str, csv_name: str, answer: str, limit: int = 0, only_reverify: bool = False, others: bool = False, approval_sol: str = "") -> str | None:
    """Returns an error text, or None when started."""
    if kind not in JOBS:
        return "unknown job"
    direct_approval = kind in {"approve", "claim_approve", "isclaim_approve"} and bool(approval_sol)
    if approval_sol and (not direct_approval or not re.fullmatch(r"[0-9]{4}", approval_sol)):
        return "enter the four-digit branch SOL for application or regular IS claim approval"
    if direct_approval:
        csv_name = (f"branches/{approval_sol}/isclaim/{approval_sol}_isclaim.csv"
                    if kind == "isclaim_approve" else f"branches/{approval_sol}/{approval_sol}.csv")
    elif csv_name not in input_csvs():
        return "pick a CSV"
    if kind != "entry" and answer != "yes":
        return "tick the confirmation"
    want, got = JOB_SCHEME.get(kind, "regular"), branches.scheme_of(ROOT / csv_name)
    if want != got:                                   # never run a job on another scheme's work list
        return f"this job needs an {SCHEME_NAME[want]} work list (picked: {SCHEME_NAME[got]})"
    other = busy_elsewhere(csv_name, kind)
    if other:
        return (f"this branch has a job running on the dashboard at localhost:{other}: "
                f"run another branch here, or wait for that job")
    with lock:
        if running():
            return "a job is already running"
        if kind in ("entry", "entry_fresh"):
            import entry_filter
            try:
                summary = entry_filter.preview(ROOT / csv_name, "old" if kind == "entry" else "fresh", only_reverify)
            except (OSError, ValueError) as exc:
                return f"cannot check the work list: {exc}"
            if not summary["ready"]:
                return "No eligible entries. Review the filtering summary and correct the missing data."
        STOP_FLAG.unlink(missing_ok=True)
        spec = dict(JOBS[kind], log=log_name(JOBS[kind]["log"]))
        args = [sys.executable, "-u", "-W", "ignore", spec["script"], csv_name]
        if kind == "entry_fresh":
            args += ["--regular-worklist"]
        if direct_approval:
            (ROOT / csv_name).parent.mkdir(parents=True, exist_ok=True)
            args += ["--branch-code", approval_sol]
        if kind != "entry":
            args += ["--yes"]                     # approval / claim jobs: confirmed in the panel
        elif only_reverify:
            args += ["--only-reverify"]           # only the Aadhaar-reverify rows, once more
        if kind not in ("entry", "entry_fresh") and limit > 0:
            args += ["--limit", str(limit)]
        if (others or direct_approval) and kind in ("claim_approve", "isclaim_approve"):
            args += ["--others"]                  # also claims filed by hand / on another computer (checked)
        env = dict(os.environ, NO_PROMPT="1", PYTHONIOENCODING="utf-8", FASALRIN_INSTANCE=INSTANCE)
        if (ROOT / spec["log"]).exists():          # the previous run's log stays readable
            os.replace(ROOT / spec["log"], ROOT / (Path(spec["log"]).stem + ".prev.log"))
        log = open(ROOT / spec["log"], "w", encoding="utf-8")
        flags = subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0
        proc = subprocess.Popen(args, cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT,
                                stdin=subprocess.DEVNULL, creationflags=flags,
                                start_new_session=os.name != "nt")   # Linux: own group, so force stop gets Chromium too
        job.clear()
        job.update(kind=kind, csv=csv_name, proc=proc, log=ROOT / spec["log"], logfh=log,
                   started=time.time(), ended=None, stop_requested=False, forced=False)
        try:                                         # other dashboards see which branch is busy here
            RUNNING_DIR.mkdir(exist_ok=True)
            (RUNNING_DIR / f"{PORT}.txt").write_text(f"{proc.pid}\n{csv_name}\n{kind}", encoding="utf-8")
        except OSError:
            pass
    return None


def stop_job(force: bool) -> None:
    with lock:
        if not running():
            return
        STOP_FLAG.touch()
        job["stop_requested"] = True
        if force:
            job["forced"] = True
            pid = job["proc"].pid
            if os.name == "nt":                   # whole tree: python + Chromium
                subprocess.run(["taskkill", "/PID", str(pid), "/T", "/F"], capture_output=True)
            else:
                import signal
                try:
                    os.killpg(pid, signal.SIGKILL)
                except (ProcessLookupError, PermissionError):
                    job["proc"].kill()


# ----------------------------------------------------------------------------
# status: parse the job log + read the CSV
# ----------------------------------------------------------------------------
ROW_RE = re.compile(r"^\[\s*(\d+)(?:/(\d+))?\]\s+(\d+)\s+(?:acct\s+(\d+)\s+)?([A-Z_:]+)\s*(.*)$")


def log_lines(path: Path) -> list[str]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return []
    return [l.rstrip() for l in re.split(r"[\r\n]+", text) if l.strip()]


def parse_log(lines: list[str]) -> dict:
    results, errors, total, alert, tally = [], [], None, "", collections.Counter()
    login_wait = False
    for l in lines:
        s = l.strip()
        m = ROW_RE.match(s)
        if m and m.group(5) not in ("working", "approving"):
            n, tot, first, acct, status, rest = m.groups()
            if tot:
                total = int(tot)
            detail = re.sub(r"\s*app=.*$|\s*\|.*$", "", rest).strip()
            app = re.search(r"app=(\d+)", rest)
            results.append({"n": int(n), "id": first, "acct": acct or first, "status": status,
                            "detail": detail, "app": app.group(1) if app else (first if acct else "")})
            tally["ERROR" if status.startswith("ERROR") else status] += 1
            login_wait = False
            continue
        if re.search(r"LOG IN|still waiting for login", s):
            login_wait = True
            alert = s.lstrip("! ")
        elif "[login] session is back" in s:
            login_wait = False
        elif s.startswith("!!"):
            alert = s.lstrip("! ")
        if s.startswith("[exc]") or "Traceback" in s or "Error:" in s and s.startswith("playwright"):
            errors.append(s[:220])
        mt = re.match(r"\[start\]\s+(\d+) rows pending", s)
        if mt:
            total = int(mt.group(1))
    info = [s for s in lines if re.match(r"\[(resume|reports|progress|approvals|list|start|stop|done|quit|search)\]"
                                          r"|\[time\] average|\[learn\]", s.strip())]
    return {"results": results, "tally": dict(tally), "total": total, "errors": errors[-10:],
            "login_wait": login_wait, "alert": alert, "info": info[-12:]}


def additional_summary(name: str) -> dict:
    """Totals for an IS additional or PRI additional work list: from its records + approvals files."""
    pri = branches.is_pri_list(ROOT / name)
    rows = (branches._pri_rows if pri else branches._additional_rows)(ROOT / name)
    b = collections.Counter(branches.additional_bucket(r.get("Status") or "") for r in rows)
    st = collections.Counter((r.get("Status") or "").strip() or "TO DO" for r in rows)
    approved = sum(1 for r in rows if r.get("Approval") == "APPROVED")
    summ = branches.hand_work_dir(ROOT / name) / f"{Path(name).stem}_summary.csv"
    return {"scheme": branches.scheme_of(ROOT / name), "rows": len(rows), "status": dict(st.most_common()),
            "reasons": branches.hand_work_counts(ROOT / name),
            "export": time.strftime("%d-%m %H:%M", time.localtime(summ.stat().st_mtime)) if summ.exists() else None,
            "additional": {"claimed": b["claimed"], "approved": approved, "on_portal": b["on_portal"], "not_found_closed": b["skipped"],
                           "mismatch": st.get("MISMATCH", 0), "check": b["check"], "errors": b["error"] - st.get("MISMATCH", 0),
                           "todo": b["todo"]}}


def isclaim_summary(name: str) -> dict:
    """Totals for an IS regular claim list (statuses from the branch's claim records + claim approvals)."""
    rows = branches._isclaim_rows(ROOT / name)
    b = collections.Counter(branches.isclaim_bucket(r.get("Status") or "") for r in rows)
    st = collections.Counter((r.get("Status") or "").strip() or "TO DO" for r in rows)
    summ = branches.hand_work_dir(ROOT / name) / f"{Path(name).stem}_summary.csv"
    approved = sum(1 for r in rows if r.get("Approval") == "APPROVED")
    return {"scheme": "isclaim", "rows": len(rows), "status": dict(st.most_common()),
            "reasons": branches.hand_work_counts(ROOT / name),
            "export": time.strftime("%d-%m %H:%M", time.localtime(summ.stat().st_mtime)) if summ.exists() else None,
            "isclaim": {"claimed": b["claimed"], "approved": approved, "to_approve": b["claimed"] - approved,
                        "waiting": b["waiting"], "hand": b["hand"], "check": b["check"], "errors": b["error"],
                        "todo": b["todo"]}}


def fresh_summary(name: str) -> dict:
    """Totals for an IS fresh work list (statuses from its progress file)."""
    rows = branches._fresh_rows(ROOT / name)
    b = collections.Counter(branches.fresh_bucket(r.get("Status") or "") for r in rows)
    st = collections.Counter((r.get("Status") or "").strip() or "TO DO" for r in rows)
    summ = branches.hand_work_dir(ROOT / name) / f"{Path(name).stem}_summary.csv"
    return {"scheme": "fresh", "rows": len(rows), "status": dict(st.most_common()),
            "reasons": branches.hand_work_counts(ROOT / name),
            "export": time.strftime("%d-%m %H:%M", time.localtime(summ.stat().st_mtime)) if summ.exists() else None,
            "fresh": {"entered": b["finished"], "hand": b["hand"], "check": b["check"], "errors": b["error"],
                      "todo": b["todo"],
                      "cc004": sum(1 for r in rows if (r.get("Scheme Code") or "").upper() == "CC004"),
                      "cc043": sum(1 for r in rows if (r.get("Scheme Code") or "").upper() == "CC043")}}


def csv_summary(name: str) -> dict:
    if branches.is_isclaim_list(ROOT / name):
        try:
            return isclaim_summary(name)
        except OSError:
            return {}
    if branches.is_fresh_list(ROOT / name):
        try:
            return fresh_summary(name)
        except OSError:
            return {}
    if branches.is_additional_list(ROOT / name) or branches.is_pri_list(ROOT / name):
        try:
            return additional_summary(name)
        except OSError:
            return {}
    try:
        with open(ROOT / name, newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
    except OSError:
        return {}
    st = collections.Counter((r.get("Status") or "").strip() or "PENDING" for r in rows)
    ap = collections.Counter((r.get("Approval") or "").strip() for r in rows if (r.get("Approval") or "").strip())
    fin = [r for r in rows if branches.f.bucket(r.get("Status") or "") == "finished"]
    # approvals: CSV column (root lists) or the branch's approvals file (branch work lists)
    sol = Path(name).parent.name if Path(name).stem == Path(name).parent.name else ""
    appr = branches.approvals(sol) if sol else {}
    approved = sum(1 for r in fin if (r.get("Approval") or "").strip() == "APPROVED"
                   or appr.get((r.get("Account No.") or "").strip()) == "APPROVED")
    b = collections.Counter(branches.f.bucket(r.get("Status") or "") for r in rows)
    claims = {"filed": 0, "approved": 0, "hand": 0, "check": 0}
    base = (ROOT / name).with_suffix("")
    for fname, key in ((f"{base}_claims.csv", "Result"), (f"{base}_claim_approvals.csv", "Result")):
        last = {}
        if os.path.exists(fname):
            with open(fname, newline="", encoding="utf-8-sig") as fh:
                for e in csv.DictReader(fh):
                    last[(e.get("Account No.") or "").strip()] = (e.get(key) or "").strip()
        for v in last.values():
            if fname.endswith("_claims.csv"):
                if v == "CLAIMED":
                    claims["filed"] += 1
                elif v.startswith("CLAIM_CHECK_PORTAL"):
                    claims["check"] += 1
                elif v in ("NO_INT_SUB_AMT", "CLAIM_ZERO_ALLOWED", "DATE_MISMATCH") or v.startswith("CLAIM_ERROR"):
                    claims["hand"] += 1
            elif v == "APPROVED":
                claims["approved"] += 1
    summ = branches.hand_work_dir(ROOT / name) / f"{Path(name).stem}_summary.csv"
    mis = collections.Counter((r.get(branches.MIS_STATUS) or "").strip().lower() for r in rows)
    return {"rows": len(rows), "status": dict(st.most_common()), "approval": dict(ap),
            "mis": {"pending": mis.get("pending", 0), "draft": mis.get("draft", 0)},
            "reasons": branches.hand_work_counts(ROOT / name), "claims": claims,
            "export": time.strftime("%d-%m %H:%M", time.localtime(summ.stat().st_mtime)) if summ.exists() else None,
            "buckets": {"finished": b["finished"], "approved": approved, "submitted": b["finished"] - approved,
                        "hand": b["hand"], "check": b["check"], "todo": b["todo"]}}


def status() -> dict:
    with lock:
        j = dict(job) if job else {}
    out = {"csvs": input_csvs(), "job": None, "now": time.time()}
    if j:
        proc = j["proc"]
        code = proc.poll()
        if code is not None and not j.get("ended"):
            job["ended"] = j["ended"] = time.time()
            try:
                j["logfh"].close()
            except Exception:
                pass
        p = parse_log(log_lines(j["log"]))
        if code is None:
            state = "login" if p["login_wait"] else ("stopping" if j["stop_requested"] else "running")
        elif j["forced"]:
            state = "stopped"
        else:
            state = "done" if code == 0 else "failed"
        out["job"] = {"kind": j["kind"], "csv": j["csv"], "state": state, "exit": code,
                      "login": JOBS[j["kind"]]["login"], "started": j["started"], "ended": j.get("ended"),
                      **p, "results": p["results"][-40:][::-1], "done": len(p["results"])}
    sel = j.get("csv") if j else None             # nothing preselected: the page asks for a list
    out["csv"] = sel
    out["summary"] = csv_summary(sel) if sel else {}
    return out


# ----------------------------------------------------------------------------
# http
# ----------------------------------------------------------------------------
SAMPLES = {"fresh.csv": "is_fresh_sample.csv",              # /sample/<key> -> samples/<file> (fixed list only)
           "regular.xlsx": "is_regular_master_sample.xlsx",
           "isclaim.xlsx": "is_regular_claim_sample.xlsx", "additional.xlsx": "is_additional_claim_sample.xlsx",
           "pri.xlsx": "pri_additional_claim_sample.xlsx", "prireg.xlsx": "pri_regular_claim_sample.xlsx"}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="application/json"):
        data = body if isinstance(body, bytes) else json.dumps(body).encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            self._send(200, (ROOT / "portal.html").read_bytes(), "text/html; charset=utf-8")
        elif self.path.startswith(("/sample/", "/samples/")):
            path = self.path.partition("?")[0]
            sample = path[len("/samples/"):] if path.startswith("/samples/") else SAMPLES.get(path[len("/sample/"):])
            if sample not in SAMPLES.values():
                return self._send(404, {"error": "sample not found"})
            if not sample or not (ROOT / "samples" / sample).exists():
                return self._send(404, {"error": "sample file missing: restart the panel to update it"})
            data = (ROOT / "samples" / sample).read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "text/csv; charset=utf-8" if sample.endswith(".csv") else
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{sample}"')
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        elif self.path.startswith("/download?"):
            from urllib.parse import unquote
            q = dict(x.split("=", 1) for x in self.path.partition("?")[2].split("&") if "=" in x)
            name = unquote(q.get("csv", ""))
            if name not in input_csvs():
                return self._send(404, {"error": "unknown list"})
            data = branches.export_zip(ROOT / name)
            if len(data) < 100:
                return self._send(404, {"error": "export it first"})
            fname = f"{Path(name).stem}_hand_work.zip"
            self.send_response(200)
            self.send_header("Content-Type", "application/zip")
            self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        elif self.path.startswith("/download_analysis"):     # one list of the IS additional report analysis
            from urllib.parse import unquote
            q = {k: unquote(v) for k, v in (x.split("=", 1) for x in self.path.partition("?")[2].split("&") if "=" in x)}
            name, fname = q.get("csv", ""), os.path.basename(q.get("file", ""))
            if name not in input_csvs() or not branches.is_analysis_list(ROOT / name):
                return self._send(404, {"error": "unknown list"})
            p = branches.analysis_dir(ROOT / name) / fname
            if not fname.startswith(Path(name).stem + "_") or not fname.endswith(".csv") or not p.exists():
                return self._send(404, {"error": "press Analyse first"})
            data = p.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "text/csv; charset=utf-8")
            self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        elif self.path.startswith("/api/masters"):
            from urllib.parse import unquote
            q = {k: unquote(v) for k, v in (x.split("=", 1) for x in self.path.partition("?")[2].split("&") if "=" in x)}
            kind = q.get("type") if q.get("type") in MASTER_KINDS else "regular"
            ms = masters(kind)
            sel = q.get("master", "") or (ms[0] if ms else "")
            if sel not in ms:
                sel = ms[0] if ms else ""
            try:
                d = MASTER_DIRS[kind]
                sols = SOLS_OF[kind](d / sel) if sel else []
                err = None
            except Exception as e:
                sols, err = [], str(e)[:200]
            self._send(200, {"type": kind, "masters": ms, "master": sel, "sols": sols, "error": err})
        elif self.path.startswith("/api/reports"):
            from urllib.parse import unquote
            q = {k: unquote(v) for k, v in (x.split("=", 1) for x in self.path.partition("?")[2].split("&") if "=" in x)}
            name = q.get("csv", "")
            if name not in input_csvs() or not name.startswith("branches/"):
                return self._send(400, {"error": "pick a branch work list"})
            self._send(200, {"csv": name, "wants": branches.wants_kind(ROOT / name),
                             "folder": branches.reports_dir(ROOT / name).relative_to(ROOT).as_posix(),
                             "reports": branches.list_reports(ROOT / name)})
        elif self.path.startswith("/api/status"):
            from urllib.parse import unquote
            q = {k: unquote(v) for k, v in (x.split("=", 1) for x in self.path.partition("?")[2].split("&") if "=" in x)}
            st = status()
            if q.get("csv") in st["csvs"] and not (st["job"] and st["job"]["state"] in ("running", "login", "stopping")):
                st["csv"], st["summary"] = q["csv"], csv_summary(q["csv"])
            if q.get("entry_kind") in ("entry", "entry_fresh") and q.get("csv") in st["csvs"]:
                import entry_filter
                name, kind = q["csv"], q["entry_kind"]
                if branches.scheme_of(ROOT / name) == "regular":
                    try:
                        st["entry_preview"] = {"csv": name, "kind": kind, "only_reverify": q.get("only_reverify") == "1",
                            **entry_filter.preview(ROOT / name, "old" if kind == "entry" else "fresh", q.get("only_reverify") == "1")}
                    except (OSError, ValueError) as exc:
                        st["entry_preview"] = {"csv": name, "kind": kind, "error": str(exc)}
            self._send(200, st)
        else:
            self._send(404, {"error": "not found"})

    def do_POST(self):
        # same-origin only: the page is served from this host; refuse cross-site form posts
        origin = self.headers.get("Origin", "")
        if origin and origin not in (f"http://{HOST}:{PORT}", f"http://localhost:{PORT}"):
            return self._send(403, {"error": "forbidden"})
        size = int(self.headers.get("Content-Length") or 0)
        if size > MAX_UPLOAD * 2:
            return self._send(413, {"error": "file too large"})
        try:
            body = json.loads(self.rfile.read(size) or b"{}")
        except ValueError:
            return self._send(400, {"error": "bad json"})
        if self.path == "/api/upload":
            err, saved = save_upload(body.get("name", ""), body.get("text", ""))
            return self._send(400 if err else 200, {"error": err} if err else {"ok": True, "csv": saved})
        if self.path == "/api/export":
            name = body.get("csv", "")
            if name not in input_csvs():
                return self._send(400, {"error": "pick a list"})
            try:
                files = branches.export_reason_csvs(ROOT / name)
            except PermissionError:
                return self._send(400, {"error": "a hand-work CSV is open in Excel — close it and export again"})
            return self._send(200, {"ok": True, "files": len(files) - 1,
                                    "folder": branches.hand_work_dir(ROOT / name).relative_to(ROOT).as_posix()})
        if self.path == "/api/corrections":
            name = body.get("csv", "")
            if name not in input_csvs():
                return self._send(400, {"error": "pick a list"})
            if running():
                return self._send(400, {"error": "stop the running job first"})
            try:
                with lock:
                    text = body.get("text", "")
                    if body.get("b64"):                 # an edited .xlsx: read it as the same table (all text)
                        text = xlsx_to_csv_text(body["b64"])
                    res = branches.apply_corrections(ROOT / name, text)
            except (ValueError, KeyError, csv.Error) as e:
                return self._send(400, {"error": str(e)[:200]})
            return self._send(200, {"ok": True, **res})
        if self.path in ("/api/add_analysis", "/api/add_claim_these", "/api/add_queue_clear"):
            name = body.get("csv", "")
            if name not in input_csvs() or not branches.is_analysis_list(ROOT / name):
                return self._send(400, {"error": "pick an IS additional or IS claim work list"})
            if self.path != "/api/add_analysis" and running() and job.get("csv") == name:
                return self._send(400, {"error": "a job is running on this work list: stop it first"})
            try:
                with lock:
                    if self.path == "/api/add_analysis":
                        res = branches.analyse_additional(ROOT / name, body.get("approved", ""), body.get("pending", ""))
                    elif self.path == "/api/add_claim_these":
                        res = branches.apply_additional_analysis(ROOT / name)
                    else:
                        branches.clear_additional_queue(ROOT / name)
                        res = {"queue": {}}
            except PermissionError:
                return self._send(400, {"error": "a list is open in Excel - close it and try again"})
            except (ValueError, KeyError) as e:
                return self._send(400, {"error": str(e)[:200]})
            return self._send(200, {"ok": True, **res})
        if self.path == "/api/master_upload":
            kind = body.get("type") if body.get("type") in MASTER_KINDS else "regular"
            err, saved = save_master(body.get("name", ""), body.get("b64", ""), kind)
            return self._send(400 if err else 200, {"error": err} if err else {"ok": True, "master": saved, "type": kind})
        if self.path == "/api/fresh_upload":         # IS fresh: the branch's own file -> its work list directly
            if running():
                return self._send(400, {"error": "stop the running job first"})
            name, b64 = body.get("name", ""), body.get("b64", "")
            if name.lower().endswith(".csv"):            # a CSV is kept as an .xlsx (all cells text)
                try:
                    name, b64 = csv_to_xlsx(name, b64)
                except ValueError as e:
                    return self._send(400, {"error": str(e)[:200]})
            err, saved = save_master(name, b64, "fresh")
            if err:
                return self._send(400, {"error": err})
            path = MASTER_DIRS["fresh"] / saved
            sols = branches.sols_fresh(path)
            if len(sols) != 1:
                path.unlink(missing_ok=True)
                return self._send(400, {"error": f"upload one branch's file: this file has {len(sols)} SOLs "
                                                 f"({', '.join(s['sol'] for s in sols)[:80]})"})
            try:
                with lock:
                    res = branches.build_fresh(path, sols[0]["sol"])
            except Exception as e:
                return self._send(400, {"error": str(e)[:200]})
            return self._send(200, {"ok": True, "type": "fresh", "file": saved, **res})
        if self.path == "/api/build":
            kind = body.get("type") if body.get("type") in MASTER_KINDS else "regular"
            m, sol = body.get("master", ""), str(body.get("sol", "")).strip()
            if m not in masters(kind) or not re.fullmatch(r"\d{3,6}", sol):
                return self._send(400, {"error": "pick a master and a SOL"})
            if running():
                return self._send(400, {"error": "stop the running job first"})
            try:
                with lock:
                    d = MASTER_DIRS[kind]
                    res = BUILD_OF[kind](d / m, sol)
            except Exception as e:
                return self._send(400, {"error": str(e)[:200]})
            return self._send(200, {"ok": True, "type": kind, **res})
        if self.path == "/api/start":
            try:
                limit = max(0, int(body.get("limit") or 0))
            except (TypeError, ValueError):
                limit = 0
            err = start_job(body.get("kind", ""), body.get("csv", ""), body.get("answer", ""), limit,
                            bool(body.get("only_reverify")),
                            bool(body.get("others")), str(body.get("approval_sol") or "").strip())
            return self._send(400 if err else 200, {"error": err} if err else {"ok": True})
        if self.path in ("/api/report_upload", "/api/report_match"):
            name = body.get("csv", "")
            if name not in input_csvs() or not name.startswith("branches/"):
                return self._send(400, {"error": "pick a branch work list"})
            if self.path == "/api/report_upload":
                import base64
                try:
                    saved = branches.save_report(ROOT / name, body.get("name", ""), base64.b64decode(body.get("b64", ""), validate=True))
                except (ValueError, OSError) as e:
                    return self._send(400, {"error": str(e)[:240]})
                try:                                  # what Match would do, before anything is written
                    with lock:
                        info = branches.report_info(branches.reports_dir(ROOT / name) / saved)
                        pending = (branches.is_analysis_list(ROOT / name)        # a pending report: no Match preview
                                   and info["kind"] != branches.wants_kind(ROOT / name))
                        preview = None if pending else branches.match_report(ROOT / name, saved, dry_run=True)
                except (ValueError, KeyError, OSError) as e:
                    return self._send(400, {"error": str(e)[:240]})
                return self._send(200, {"ok": True, "report": saved, "info": info, "preview": preview,
                                        "pending": pending})
            if running() and job.get("csv") == name:
                return self._send(400, {"error": "a job is running on this work list: stop it first, then match"})
            rp = branches.reports_dir(ROOT / name) / os.path.basename(body.get("report", ""))
            if (branches.is_analysis_list(ROOT / name) and rp.exists()
                    and branches.report_info(rp)["kind"] != branches.wants_kind(ROOT / name)):
                return self._send(400, {"error": "this is the pending report: use it in Report analysis below"})
            try:
                with lock:
                    res = branches.match_report(ROOT / name, body.get("report", ""))
            except PermissionError:
                return self._send(400, {"error": "the work list is open in Excel - close it and match again"})
            except (ValueError, KeyError, OSError) as e:
                return self._send(400, {"error": str(e)[:240]})
            return self._send(200, {"ok": True, **res})
        if self.path == "/api/reset_hand":
            name = body.get("csv", "")
            if name not in input_csvs():
                return self._send(400, {"error": "pick a list"})
            if running() and job.get("csv") == name:
                return self._send(400, {"error": "a job is running on this work list: stop it first"})
            try:
                with lock:
                    res = branches.reset_hand_work(ROOT / name, list(body.get("reasons") or []))
            except PermissionError:
                return self._send(400, {"error": "the work list is open in Excel - close it and reset again"})
            except (ValueError, KeyError, OSError) as e:
                return self._send(400, {"error": str(e)[:200]})
            return self._send(200, {"ok": True, **res})
        if self.path == "/api/stop":
            stop_job(bool(body.get("force")))
            return self._send(200, {"ok": True})
        self._send(404, {"error": "not found"})


class Server(ThreadingHTTPServer):
    # Windows lets a second panel bind the same port (SO_REUSEADDR) and requests then hang between the two:
    # claim the port exclusively, so a second start fails and just opens the panel that already runs
    allow_reuse_address = os.name != "nt"

    def server_bind(self):
        if os.name == "nt":
            import socket
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()


def main():
    url = f"http://localhost:{PORT}"
    try:
        srv = Server((HOST, PORT), Handler)
    except OSError:
        print(f"The Fasalrin panel is already running at {url} - opening it. (Only one panel can run at a time.)")
        if "--no-browser" not in sys.argv:
            webbrowser.open(url)
        return
    print(f"Fasalrin script: {url}   (Ctrl+C to quit)")
    if "--no-browser" not in sys.argv:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nbye")


if __name__ == "__main__":
    main()
