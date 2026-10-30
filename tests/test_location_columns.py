import unittest
from unittest.mock import Mock, patch
from pathlib import Path
import csv
from openpyxl import load_workbook
import branches

import fasalrin_fresh as fresh
import fasalrin_regular as regular


class LocationColumnTests(unittest.TestCase):
    def test_explicit_taluka_overrides_address_score(self):
        page, scope = Mock(), Mock()
        with patch.object(fresh, "wait_for"), patch.object(fresh, "options_of", return_value=["Select", "Dahod", "Devgadhbaria"]), \
             patch.object(fresh, "select_label") as select, \
             patch.object(fresh, "form_values", return_value={"resSubDistrictId": "Devgadhbaria"}):
            chosen = fresh.pick_block(page, scope, "resSubDistrictId", "Dabhva", "DIST DAHOD", taluka="Devgadbaria")
        self.assertEqual(chosen, "Devgadhbaria")
        select.assert_called_once_with(scope, "resSubDistrictId", "Devgadhbaria", page)

    def test_unknown_explicit_taluka_does_not_fall_back_to_dahod(self):
        with patch.object(fresh, "wait_for"), patch.object(fresh, "options_of", return_value=["Select", "Dahod"]):
            with self.assertRaisesRegex(RuntimeError, "TALUKA"):
                fresh.pick_block(Mock(), Mock(), "resSubDistrictId", "Dabhva", "DIST DAHOD", taluka="Devgadbaria")

    def test_explicit_taluka_must_stick_in_portal(self):
        with patch.object(fresh, "wait_for"), patch.object(fresh, "options_of", return_value=["Select", "Dahod", "Devgadhbaria"]), \
             patch.object(fresh, "select_label"), patch.object(fresh, "form_values", return_value={"resSubDistrictId": "Dahod"}):
            with self.assertRaisesRegex(RuntimeError, "portal taluka"):
                fresh.pick_block(Mock(), Mock(), "resSubDistrictId", "Dabhva", "DIST DAHOD", taluka="Devgadbaria")

    def test_master_location_spelling_variants(self):
        examples = [
            ("Lunawada", "Lunavada"),
            ("Dohad", "Dahod"),
            ("Devgadhbarida", "Devgadhbaria"),
            ("Devgadbaria", "Devgadhbaria"),
        ]
        for source, portal in examples:
            with self.subTest(source=source):
                self.assertEqual(fresh.master_location_option(source, ["Select", portal], "TALUKA"), portal)

    def test_regular_row_carries_district_and_taluka(self):
        header = ["ADDRESS", "Branch Name", "DISTRICT", "TALUKA"]
        row = ["DIST DAHOD", "Dabhva", "Dahod", "Devgadbaria"]
        self.assertEqual(regular.land_of(row, dict(zip(header, range(len(header)))))["taluka"], "Devgadbaria")
        self.assertEqual(regular.land_of(row, dict(zip(header, range(len(header)))))["district"], "Dahod")

    def test_master_templates_end_with_location_columns(self):
        root = Path(__file__).resolve().parents[1]
        wb = load_workbook(root / "samples/is_regular_master_sample.xlsx", read_only=True)
        self.assertEqual([c.value for c in next(wb.worksheets[0].iter_rows())][-2:], ["DISTRICT", "TALUKA"])
        wb.close()
        with (root / "samples/is_fresh_sample.csv").open(newline="") as fh:
            self.assertEqual(next(csv.reader(fh))[-2:], ["DISTRICT", "TALUKA"])

    def test_blank_location_columns_are_kept_for_work_list_header(self):
        root = Path(__file__).resolve().parents[1]
        with patch.object(branches, "_cache", {}):
            row = branches.read_master(root / "samples/is_regular_master_sample.xlsx")[0]
        self.assertIn("DISTRICT", row)
        self.assertIn("TALUKA", row)
        self.assertEqual(row["TALUKA"], "")


if __name__ == "__main__":
    unittest.main()
