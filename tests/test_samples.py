import threading
import unittest
import urllib.request
import urllib.error
from pathlib import Path
from openpyxl import load_workbook
import branches
import portal


class SampleTests(unittest.TestCase):
    def test_loan_master_template_does_not_include_portal_column(self):
        sample = portal.ROOT / 'samples/is_regular_master_sample.xlsx'
        wb = load_workbook(sample, read_only=True)
        headers = [c.value for c in next(wb.worksheets[0].iter_rows())]
        wb.close()
        self.assertNotIn('PORTAL', headers)
        self.assertEqual(headers[-2:], ['DISTRICT', 'TALUKA'])
        self.assertNotIn('PORTAL', branches.read_master(sample)[0])

    def test_claim_templates_have_required_columns(self):
        for name, required in [('is_regular_claim_sample.xlsx', branches.ISCLAIM_REQUIRED),
                               ('is_additional_claim_sample.xlsx', branches.ADDITIONAL_REQUIRED),
                               ('pri_additional_claim_sample.xlsx', branches.PRI_REQUIRED),
                               ('pri_regular_claim_sample.xlsx', branches.PRIREG_REQUIRED)]:
            with self.subTest(name=name):
                wb = load_workbook(portal.ROOT / 'samples' / name)
                sheet = wb['Sheet1']
                self.assertTrue(set(required).issubset({c.value for c in sheet[1]}))
                self.assertEqual(sheet['C2'].number_format, '@')
                self.assertIsNone(sheet['C2'].value)
                wb.close()

    def test_static_sample_downloads_and_allowlist(self):
        server = portal.Server(('127.0.0.1', 0), portal.Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f'http://127.0.0.1:{server.server_port}'
        try:
            for filename in portal.SAMPLES.values():
                with urllib.request.urlopen(base + '/samples/' + filename) as r:
                    self.assertEqual(r.read(), (portal.ROOT / 'samples' / filename).read_bytes())
                    self.assertIn(filename, r.headers['Content-Disposition'])
            for path in ('/samples/../portal.py', '/samples/nope.xlsx'):
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(base + path)
                self.assertEqual(error.exception.code, 404)
        finally:
            server.shutdown()
            server.server_close()
            thread.join()
